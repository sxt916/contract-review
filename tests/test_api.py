from pathlib import Path
from html.parser import HTMLParser
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from contract_review.config import get_settings
from contract_review.db import init_db
from contract_review.main import app
from contract_review.worker import run_once


class _NavLinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_nav = False
        self.current_href = None
        self.current_text = []
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "nav":
            self.in_nav = True
        elif self.in_nav and tag == "a":
            self.current_href = dict(attrs).get("href")
            self.current_text = []

    def handle_data(self, data):
        if self.current_href is not None:
            self.current_text.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.current_href is not None:
            self.links.append(("".join(self.current_text).strip(), self.current_href))
            self.current_href = None
        elif tag == "nav":
            self.in_nav = False


class _PageLinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.current_href = None
        self.current_text = []
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.current_href = dict(attrs).get("href")
            self.current_text = []

    def handle_data(self, data):
        if self.current_href is not None:
            self.current_text.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.current_href is not None:
            self.links.append((" ".join("".join(self.current_text).split()), self.current_href))
            self.current_href = None


def test_home_page_offers_both_contract_modules():
    with TestClient(app) as client:
        response = client.get("/")

    assert response.status_code == 200
    parser = _PageLinkParser()
    parser.feed(response.text)
    links = {href: text for text, href in parser.links}
    assert set(links) == {"comparison.html", "amount-review.html"}
    assert "合同对比" in links["comparison.html"]
    assert "金额审核" in links["amount-review.html"]


def test_work_pages_offer_comparison_and_amount_review_navigation():
    expected = [("合同对比", "comparison.html"), ("金额审核", "amount-review.html")]
    with TestClient(app) as client:
        for path in (
            "/comparison.html",
            "/comparison-results.html",
            "/amount-review.html",
            "/amount-review-results.html",
        ):
            response = client.get(path)
            assert response.status_code == 200
            parser = _NavLinkParser()
            parser.feed(response.text)
            assert parser.links == expected, path


def test_amount_upload_worker_and_delete(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    init_db()
    path = tmp_path / "no-amount-relationships.docx"
    other_path = tmp_path / "other-contract.docx"
    Document().save(path)
    Document().save(other_path)
    with TestClient(app) as client, path.open("rb") as stream:
        response = client.post("/api/amount-reviews", files={"files": (path.name, stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert response.status_code == 202
        task_id = response.json()["tasks"][0]["id"]
        assert run_once()
        task = client.get(f"/api/amount-reviews/{task_id}").json()
        assert task["status"] == "passed"
        with other_path.open("rb") as other_stream:
            other_response = client.post("/api/amount-reviews", files={"files": (other_path.name, other_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        other_task_id = other_response.json()["tasks"][0]["id"]
        assert run_once()
        current_batch = client.get("/api/amount-reviews", params={"task_ids": task_id}).json()
        assert current_batch["total"] == 1
        assert current_batch["items"][0]["id"] == task_id
        filtered = client.get("/api/amount-reviews", params={"filename": "no-amount", "status": "passed"}).json()
        assert filtered["total"] == 1
        assert client.get("/api/amount-reviews", params={"task_ids": "invalid"}).status_code == 400
        assert task["temp_path"] is None
        assert client.delete(f"/api/amount-reviews/{task_id}").status_code == 204
        assert client.delete(f"/api/amount-reviews/{other_task_id}").status_code == 204
    get_settings.cache_clear()


def test_amount_upload_keeps_valid_docx_and_reports_each_rejected_file(tmp_path, monkeypatch):
    """One invalid upload must not discard valid contracts from the same batch."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    init_db()
    valid = BytesIO()
    Document().save(valid)

    with TestClient(app) as client:
        response = client.post(
            "/api/amount-reviews",
            files=[
                ("files", ("有效合同.docx", valid.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", ("报价单.pdf", b"%PDF-invalid", "application/pdf")),
                ("files", ("损坏合同.docx", b"not-a-word-file", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ],
        )

        assert response.status_code == 202
        body = response.json()
        assert [task["filename"] for task in body["tasks"]] == ["有效合同.docx"]
        assert [item["filename"] for item in body["rejected"]] == ["报价单.pdf", "损坏合同.docx"]
        assert "仅支持 .docx" in body["rejected"][0]["reason"]
        assert "不是有效的 DOCX" in body["rejected"][1]["reason"]
        assert client.delete(f"/api/amount-reviews/{body['tasks'][0]['id']}").status_code == 204
    get_settings.cache_clear()


def test_amount_upload_accepts_first_ten_valid_files_and_reports_the_rest(tmp_path, monkeypatch):
    """The eleventh valid contract must be reported instead of failing the whole batch."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    init_db()
    valid = BytesIO()
    Document().save(valid)
    files = [
        ("files", (f"合同{index}.docx", valid.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        for index in range(1, 12)
    ]

    with TestClient(app) as client:
        response = client.post("/api/amount-reviews", files=files)

        assert response.status_code == 202
        body = response.json()
        assert len(body["tasks"]) == 10
        assert body["rejected"] == [{"filename": "合同11.docx", "reason": "超过一次最多 10 份的限制"}]
        for task in body["tasks"]:
            assert client.delete(f"/api/amount-reviews/{task['id']}").status_code == 204
    get_settings.cache_clear()


def test_amount_upload_silently_ignores_hidden_and_word_temporary_files(tmp_path, monkeypatch):
    """Finder metadata and Word lock files must not become tasks or rejections."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    init_db()
    valid = BytesIO()
    Document().save(valid)

    with TestClient(app) as client:
        response = client.post(
            "/api/amount-reviews",
            files=[
                ("files", ("有效合同.docx", valid.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", (".DS_Store", b"finder", "application/octet-stream")),
                ("files", ("._合同副本.docx", b"apple-double", "application/octet-stream")),
                ("files", ("~$有效合同.docx", b"word-lock", "application/octet-stream")),
            ],
        )

        assert response.status_code == 202
        body = response.json()
        assert [task["filename"] for task in body["tasks"]] == ["有效合同.docx"]
        assert body["rejected"] == []
        assert client.delete(f"/api/amount-reviews/{body['tasks'][0]['id']}").status_code == 204
    get_settings.cache_clear()


def test_pair_preview_api(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()


def test_comparison_mock_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("COMPARISON_ADAPTER", "mock")
    get_settings.cache_clear()
    init_db()
    word = tmp_path / "NPA2026-XWZCG-1874.docx"
    pdf = tmp_path / "NPA2026-XWZCG-1874-盖章.pdf"
    Document().save(word)
    writer = PdfWriter(); writer.add_blank_page(width=595, height=842)
    with pdf.open("wb") as output:
        writer.write(output)
    with TestClient(app) as client, word.open("rb") as word_stream, pdf.open("rb") as pdf_stream:
        response = client.post(
            "/api/comparisons",
            files=[
                ("word_files", (word.name, word_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("pdf_files", (pdf.name, pdf_stream, "application/pdf")),
            ],
            data={
                "pairs_json": '[{"word_index":0,"pdf_index":0,"contract_no":"NPA2026-XWZCG-1874"}]',
                "ignore_options_json": '{"comments":true,"headerfooter":false,"stamp":true,"symbols":false,"watermark":false}',
            },
        )
        assert response.status_code == 202
        task_id = response.json()["task_ids"][0]
        assert run_once()
        task = client.get("/api/comparisons").json()["items"][0]
        assert task["status"] == "processing"
        assert task["textin_task_id"]
        assert '"comments": true' in task["ignore_options_json"]
        assert run_once()
        task = client.get("/api/comparisons").json()["items"][0]
        assert task["status"] == "completed"
        filtered = client.get("/api/comparisons", params={"filename": "XWZCG-1874", "status": "completed"}).json()
        assert filtered["total"] == 1
        assert filtered["items"][0]["id"] == task_id
        current_batch = client.get("/api/comparisons", params={"task_ids": task_id}).json()
        assert current_batch["total"] == 1
        assert client.get("/api/comparisons", params={"status": "unknown"}).status_code == 400
        preview = client.get(f"/api/comparisons/{task_id}/preview")
        assert preview.status_code == 200
        assert "/api/mock-preview/" in preview.json()["preview_url"]
        assert client.delete(f"/api/comparisons/{task_id}").status_code == 204
    get_settings.cache_clear()
    with TestClient(app) as client:
        response = client.post("/api/comparisons/pair-preview", json={"word_names": ["NPA2026-XWZCG-1874.docx"], "pdf_names": ["NPA2026-XWZCG-1874-盖章.pdf"]})
        assert response.status_code == 200
        assert len(response.json()["pairs"]) == 1
    get_settings.cache_clear()
