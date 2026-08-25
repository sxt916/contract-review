from pathlib import Path

from docx import Document

from contract_review.reviewer import review_docx


def make_contract(path: Path, wrong_line: bool = False, payment: str | list[str] | None = None):
    doc = Document()
    doc.add_paragraph("合同编号：NPA2026-XWZCG-1874")
    doc.add_paragraph("第一条 合同标的")
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["名称", "数量", "单价（元）", "总价（元）"]):
        cell.text = value
    values = [["商品A", "2", "10.005", "20.00" if wrong_line else "20.01"], ["商品B", "1", "79.99", "79.99"]]
    for values_row in values:
        cells = table.add_row().cells
        for cell, value in zip(cells, values_row):
            cell.text = value
    cells = table.add_row().cells
    for cell, value in zip(cells, ["合计", "", "", "100.00"]):
        cell.text = value
    doc.add_paragraph("第八条 结算方式及期限")
    payment_paragraphs = payment if isinstance(payment, list) else [payment or "1、验收后支付100%，即人民币100.00元（人民币壹佰元整）。"]
    for paragraph in payment_paragraphs:
        doc.add_paragraph(paragraph)
    doc.add_paragraph("第九条 违约责任")
    doc.add_paragraph("违约金为合同金额的10%。")
    doc.save(path)


def test_valid_contract(tmp_path):
    path = tmp_path / "contract.docx"
    make_contract(path)
    result = review_docx(path)
    assert result["status"] == "passed"
    assert result["contract_no"] == "NPA2026-XWZCG-1874"


def test_wrong_line_and_total(tmp_path):
    path = tmp_path / "contract.docx"
    make_contract(path, wrong_line=True)
    result = review_docx(path)
    types = {error["error_type"] for error in result["errors"]}
    assert "明细计算不一致" in types
    assert "明细合计不一致" in types


def test_fixed_and_ratio_mismatch(tmp_path):
    path = tmp_path / "contract.docx"
    make_contract(path, payment="1、预付款50%，即40.00元；2、验收款60.00元。")
    result = review_docx(path)
    assert any(e["error_type"] == "付款比例与金额不一致" for e in result["errors"])


def test_tax_and_tolerance_percentages_are_not_payment_ratios(tmp_path):
    """Reintroducing first-percentage selection would turn 13% tax into payment."""
    path = tmp_path / "contract.docx"
    make_contract(
        path,
        payment=[
            "产品应以实际交货数量结算，容差率为0%。",
            "乙方开具13%全额增值税专用发票后，甲方90天内支付合同总额的100%。",
        ],
    )

    result = review_docx(path)

    assert result["status"] == "passed"
    assert result["errors"] == []


def test_multiple_real_payment_ratios_in_one_paragraph_are_all_counted(tmp_path):
    """Dropping later real payment clauses would make a valid split payment fail."""
    path = tmp_path / "contract.docx"
    make_contract(path, payment="预付款20%，验收合格后支付80%。")

    result = review_docx(path)

    assert result["status"] == "passed"


def test_missing_optional_amount_sections_are_skipped(tmp_path):
    """Removing the optional table/payment checks must not make a contract fail."""
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("本合同未约定可计算的金额关系。")
    doc.save(path)

    result = review_docx(path)

    assert result["status"] == "passed"
    assert result["errors"] == []


def test_amount_only_table_checks_item_sum(tmp_path):
    """Removing sum validation from tables without quantity/unit price must fail."""
    path = tmp_path / "contract.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "金额"
    for name, total in [("服务A", "100.00"), ("服务B", "200.00")]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = total
    cells = table.add_row().cells
    cells[0].text = "合计"
    cells[1].text = "301.00"
    doc.save(path)

    result = review_docx(path)

    sum_error = next(error for error in result["errors"] if error["error_type"] == "明细合计不一致")
    assert "100.00 + 200.00 = 300.00" in sum_error["original"]


def test_weight_column_is_inferred_when_eighty_percent_of_rows_match(tmp_path):
    """Removing validated weight inference would miss the one incorrect line total."""
    path = tmp_path / "contract.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["名称", "理论重量（KG）", "含税单价（元）", "金额（元）"]):
        cell.text = value
    rows = [
        ("货物A", "10", "10", "100"),
        ("货物B", "20", "5", "100"),
        ("货物C", "30", "2", "60"),
        ("货物D", "40", "3", "120"),
        ("货物E", "50", "4", "201"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    cells = table.add_row().cells
    for cell, value in zip(cells, ["合计", "", "", "581"]):
        cell.text = value
    doc.save(path)

    result = review_docx(path)

    assert any(error["error_type"] == "明细计算不一致" for error in result["errors"])
    assert any("理论重量" in check for check in result["summary"]["checks"])


def test_total_row_uses_currency_amount_instead_of_tax_rate(tmp_path):
    """Treating the nearby 13% tax rate as the contract total must fail."""
    path = tmp_path / "contract.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["名称", "数量", "单价（元）", "总价（元）"]):
        cell.text = value
    cells = table.add_row().cells
    for cell, value in zip(cells, ["货物", "1", "4289.00", "4289.00"]):
        cell.text = value
    cells = table.add_row().cells
    cells[0].text = "合计"
    cells[3].text = "人民币【4,289.00】元，大写【肆仟贰佰捌拾玖元整】（含【13】%增值税）"
    doc.save(path)

    result = review_docx(path)

    assert result["status"] == "passed"
    assert result["summary"]["contract_total"] == "4289.00"


def test_standalone_uppercase_mismatch_is_reported(tmp_path):
    """Removing standalone Arabic/uppercase comparison must fail."""
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("合同总额人民币【4,289.00】元，大写【肆仟元整】。")
    doc.save(path)

    result = review_docx(path)

    assert any(error["error_type"] == "中文大写金额不一致" for error in result["errors"])


def test_merged_total_row_text_is_not_repeated_in_error(tmp_path):
    """Reading a merged cell through every grid cell must not duplicate its text."""
    path = tmp_path / "contract.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["名称", "数量", "单价", "总价"]):
        cell.text = value
    cells = table.add_row().cells
    for cell, value in zip(cells, ["货物", "1", "4289", "4289"]):
        cell.text = value
    cells = table.add_row().cells
    cells[0].text = "合计"
    merged = cells[1].merge(cells[3])
    merged.text = "人民币【4,289.00】元，大写【肆仟贰佰捌拾捌元整】"
    doc.save(path)

    result = review_docx(path)
    error = next(error for error in result["errors"] if error["error_type"] == "中文大写金额不一致")

    assert error["original"].count("人民币") == 1


def test_identical_errors_are_deduplicated(tmp_path):
    """Repeating an identical source expression must not create duplicate cards."""
    path = tmp_path / "contract.docx"
    doc = Document()
    text = "合同总额人民币【4,289.00】元，大写【肆仟元整】。"
    doc.add_paragraph(text)
    doc.add_paragraph(text)
    doc.save(path)

    result = review_docx(path)
    uppercase_errors = [error for error in result["errors"] if error["error_type"] == "中文大写金额不一致"]

    assert len(uppercase_errors) == 1
